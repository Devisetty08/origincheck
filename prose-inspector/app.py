"""Prose Inspector: a Streamlit dashboard for spotting monotonous rhythm, stock phrasing and repetition in a draft.

Run:  streamlit run app.py
"""
from __future__ import annotations

import html
import io
import json

import pandas as pd
import streamlit as st

from prose import DEFAULT_PATTERNS, analyze

st.set_page_config(page_title="Prose Inspector", page_icon="📝", layout="wide")

CATEGORY_COLOURS = {"Stock transitions": "#0078d4", "Fillers": "#d97706", "Clichés and buzzwords": "#d13438", "Run-ups": "#9333ea"}

# ---------------------------------------------------------------- sidebar: configuration
with st.sidebar:
    st.header("Settings")
    long_words = st.slider("Flag sentences longer than (words)", 20, 60, 35)
    uniform_cv = st.slider("Flag paragraphs with spread (CV) below", 0.05, 0.6, 0.25, 0.05,
                           help="Coefficient of variation = standard deviation ÷ average sentence length. "
                                "Low values mean every sentence in the paragraph is about the same length.")
    close_window = st.slider("Repeated word within (sentences)", 1, 4, 2)
    st.divider()
    st.subheader("Phrase dictionary")
    st.caption("One phrase per line. Matching ignores case and only matches whole words.")
    if st.button("Reset to defaults"):
        for cat in DEFAULT_PATTERNS:
            st.session_state[f"pat_{cat}"] = "\n".join(DEFAULT_PATTERNS[cat])
    patterns = {}
    for cat, phrases in DEFAULT_PATTERNS.items():
        st.session_state.setdefault(f"pat_{cat}", "\n".join(phrases))
        patterns[cat] = st.text_area(cat, key=f"pat_{cat}", height=140).splitlines()
    new_cat = st.text_input("Add a category (optional)")
    if new_cat:
        patterns[new_cat] = st.text_area(f"{new_cat} phrases", key=f"pat_custom_{new_cat}").splitlines()


def read_upload(f) -> str:
    name = f.name.lower()
    if name.endswith(".docx"):
        import docx  # python-docx

        d = docx.Document(io.BytesIO(f.getvalue()))
        return "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
    return f.getvalue().decode("utf-8", errors="ignore")


# ---------------------------------------------------------------- input
st.title("Prose Inspector")
st.write("Find monotonous sentence rhythm, stock phrasing and repetition in your draft, then edit those spots yourself. "
         "This tool reports; it never rewrites your text.")

src = st.radio("Input", ["Paste text", "Upload a file"], horizontal=True, label_visibility="collapsed")
if src == "Paste text":
    text = st.text_area("Your draft", height=220, placeholder="Paste your text here…", key="draft")
else:
    up = st.file_uploader("Upload .txt, .md or .docx", type=["txt", "md", "docx"])
    text = read_upload(up) if up else ""

if not text.strip():
    st.info("Paste or upload some text to see the analysis.")
    st.stop()

r = analyze(text, patterns, long_words=long_words, uniform_cv=uniform_cv, close_window=close_window)
if not r.sentences:
    st.warning("No sentences found.")
    st.stop()

# ---------------------------------------------------------------- headline metrics
o = r.overall
c = st.columns(6)
c[0].metric("Words", f"{sum(s.words for s in r.sentences):,}")
c[1].metric("Sentences", o["sentences"])
c[2].metric("Avg sentence (words)", o["mean"])
c[3].metric("Spread, SD (words)", o["sd"], help=f"Coefficient of variation {o['cv']}")
c[4].metric("Stock phrases", len(r.hits))
c[5].metric("Long sentences", len(r.long_sentences))

tabs = st.tabs(["Rhythm", "Phrase audit", "Repetition", "Annotated text", "Export"])

# ---------------------------------------------------------------- rhythm
with tabs[0]:
    st.subheader("Sentence length, in order")
    st.caption("Long runs of bars at the same height read as monotonous. Peaks are your long sentences.")
    st.bar_chart(pd.DataFrame({"words": [s.words for s in r.sentences]},
                              index=pd.RangeIndex(1, len(r.sentences) + 1, name="sentence")), height=220)
    left, right = st.columns(2)
    with left:
        st.subheader("Distribution")
        bins = pd.cut(pd.Series([s.words for s in r.sentences]), bins=[0, 5, 10, 15, 20, 25, 30, 40, 1000],
                      labels=["1–5", "6–10", "11–15", "16–20", "21–25", "26–30", "31–40", "40+"])
        st.bar_chart(bins.value_counts(sort=False).rename("sentences"), height=220)
    with right:
        st.subheader("By paragraph")
        df = pd.DataFrame(r.paragraphs).rename(columns={"para": "¶", "mean": "avg", "sd": "SD", "cv": "CV", "first": "starts with"})
        st.dataframe(df[["¶", "sentences", "avg", "SD", "CV", "min", "max", "starts with"]], hide_index=True, width="stretch")
    if r.uniform_paragraphs:
        for p in r.uniform_paragraphs:
            st.warning(f"**Paragraph {p['para']}** — {p['sentences']} sentences all around {p['mean']} words (CV {p['cv']}). "
                       f"Consider combining two ideas into one longer sentence, or giving one point a short sentence of its own. "
                       f"Starts: “{p['first']}…”")
    else:
        st.success("No paragraph has a monotonous run of same-length sentences.")
    if r.long_sentences:
        with st.expander(f"{len(r.long_sentences)} sentence(s) longer than {long_words} words", expanded=True):
            for s in r.long_sentences:
                st.markdown(f"- **¶{s.para}, {s.words} words:** {s.text}")

# ---------------------------------------------------------------- phrase audit
with tabs[1]:
    if not r.hits:
        st.success("None of the configured phrases appear.")
    else:
        counts = pd.DataFrame(r.hits).groupby("category").size()
        mc = st.columns(max(1, len(patterns)))
        for i, cat in enumerate(patterns):
            mc[i % len(mc)].metric(cat, int(counts.get(cat, 0)))
        table = (pd.DataFrame(r.hits).groupby(["category", "phrase"])
                 .agg(count=("phrase", "size"), sentences=("sentence", lambda x: ", ".join(map(str, sorted(set(x))))))
                 .reset_index().sort_values("count", ascending=False))
        st.dataframe(table, hide_index=True, width="stretch")
        st.caption("These phrases aren't wrong. Repeated stock phrasing makes writing feel generic; check whether "
                   "each one earns its place or whether a plain statement says it better.")

# ---------------------------------------------------------------- repetition
with tabs[2]:
    left, right = st.columns(2)
    with left:
        st.subheader("Most-used words")
        if r.top_words:
            st.bar_chart(pd.Series(dict(r.top_words), name="uses"), horizontal=True, height=320)
        else:
            st.success("No content word is used three or more times.")
    with right:
        st.subheader("Repeated phrases")
        if r.top_phrases:
            st.dataframe(pd.DataFrame(r.top_phrases, columns=["phrase", "uses"]), hide_index=True, width="stretch")
        else:
            st.success("No two- or three-word phrase repeats.")
    st.subheader(f"Same word again within {close_window} sentence(s)")
    if r.close_repeats:
        st.dataframe(pd.DataFrame([{"word": x["word"], "sentences": f"{x['sentences'][0]} → {x['sentences'][1]}"} for x in r.close_repeats]),
                     hide_index=True, width="stretch")
    else:
        st.success("No close repeats.")

# ---------------------------------------------------------------- annotated text
with tabs[3]:
    legend = " ".join(f'<span style="border-bottom:3px solid {col};padding:0 4px">{html.escape(cat)}</span>' for cat, col in CATEGORY_COLOURS.items())
    st.markdown(f"{legend} &nbsp; <span style='background:rgba(234,88,12,.12);padding:0 4px'>long sentence</span>", unsafe_allow_html=True)
    # Cut the text at every phrase and long-sentence boundary, then style each piece on its own (no nesting to get wrong).
    src_text = text.replace("\r\n", "\n")
    longs = [(s.start, src_text.find(s.text, s.start) + len(s.text)) for s in r.long_sentences]
    marks = [(h["start"], h["end"], CATEGORY_COLOURS.get(h["category"], "#4a5d75"), h["category"]) for h in r.hits]
    cuts = sorted({0, len(src_text), *(x for a, b in longs for x in (a, b)), *(x for a, b, _, _ in marks for x in (a, b))})
    out = []
    for a, b in zip(cuts, cuts[1:]):
        piece = html.escape(src_text[a:b])
        style = []
        if any(s <= a and b <= e for s, e in longs):
            style.append("background:rgba(234,88,12,.12)")
        m = next((m for m in marks if m[0] <= a and b <= m[1]), None)
        if m:
            style.append(f"border-bottom:3px solid {m[2]}")
        title = f' title="{html.escape(m[3])}"' if m else ""
        out.append(f'<span style="{";".join(style)}"{title}>{piece}</span>' if style else piece)
    body = "".join(out).replace("\n", "<br>")  # blank lines inside st.markdown HTML would end the HTML block
    st.markdown(f"<div style='line-height:1.8;font-family:Georgia,serif;font-size:1.05rem'>{body}</div>",
                unsafe_allow_html=True)

# ---------------------------------------------------------------- export
with tabs[4]:
    report = {
        "overall": r.overall,
        "paragraphs": r.paragraphs,
        "uniform_paragraphs": [p["para"] for p in r.uniform_paragraphs],
        "long_sentences": [{"para": s.para, "words": s.words, "text": s.text} for s in r.long_sentences],
        "phrase_hits": [{k: h[k] for k in ("category", "phrase", "sentence", "para")} for h in r.hits],
        "top_words": r.top_words, "repeated_phrases": r.top_phrases, "close_repeats": r.close_repeats,
    }
    md = [f"# Prose Inspector report\n",
          f"- {o['sentences']} sentences, average {o['mean']} words, SD {o['sd']} (CV {o['cv']})",
          f"- {len(r.hits)} stock-phrase hits, {len(r.long_sentences)} long sentences, "
          f"{len(r.uniform_paragraphs)} monotonous paragraph(s)\n", "## Stock phrases"]
    md += [f"- {h['phrase']} ({h['category']}) — sentence {h['sentence']}" for h in r.hits] or ["- none"]
    md += ["\n## Long sentences"] + ([f"- ¶{s.para} ({s.words} words): {s.text}" for s in r.long_sentences] or ["- none"])
    md += ["\n## Most-used words"] + ([f"- {w}: {c}" for w, c in r.top_words] or ["- none"])
    a, b = st.columns(2)
    a.download_button("Download report (Markdown)", "\n".join(md), "prose_report.md", "text/markdown", width="stretch")
    b.download_button("Download report (JSON)", json.dumps(report, indent=2), "prose_report.json", "application/json", width="stretch")
