#!/usr/bin/env python3
"""
origincheck - a local, Turnitin-style similarity checker.

Compares a submission (Word, PDF, text, Markdown, HTML) against a folder of
source documents and, optionally, the web. Produces a colour-coded HTML
similarity report: overall similarity index, per-source percentages and
highlighted matching passages.

Usage:
    python origincheck.py draft.docx --sources ./references
    python origincheck.py draft.pdf --sources ./refs --exclude-quotes --exclude-bib --min-words 10
    python origincheck.py draft.docx --sources ./refs --web --web-queries 25   (needs BRAVE_API_KEY)
    python origincheck.py draft.docx --sources ./refs --save-to-repo           (keep this draft for future checks)

Run `python origincheck.py -h` for every option.
"""

import argparse
import hashlib
import html
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED = {".docx", ".pdf", ".txt", ".md", ".html", ".htm", ".rtf"}
REPO_DIR = Path(__file__).resolve().parent / "repository"

# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "nav", "footer", "header", "svg"}
    BLOCK = {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "section", "article"}

    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def html_to_text(raw: str) -> str:
    p = _HTMLText()
    p.feed(raw)
    return re.sub(r"\n\s*\n+", "\n\n", "".join(p.parts)).strip()


def rtf_to_text(raw: str) -> str:
    raw = re.sub(r"\\par[d]?", "\n", raw)
    raw = re.sub(r"\\'[0-9a-f]{2}", "", raw)
    raw = re.sub(r"\\[a-z]+-?\d* ?", "", raw)
    return re.sub(r"[{}]", "", raw)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx  # python-docx

        d = docx.Document(str(path))
        paras = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                paras.append(" | ".join(c.text for c in row.cells))
        return "\n\n".join(p for p in paras if p.strip())
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [(pg.extract_text() or "") for pg in reader.pages]
        text = "\n\n".join(pages)
        text = re.sub(r"-\n(\w)", r"\1", text)          # re-join hyphenated line breaks
        text = re.sub(r"(?<![.\n:])\n(?!\n)", " ", text)  # unwrap soft line breaks
        return text
    raw = path.read_text(encoding="utf-8", errors="ignore")
    if ext in (".html", ".htm"):
        return html_to_text(raw)
    if ext == ".rtf":
        return rtf_to_text(raw)
    return raw


# --------------------------------------------------------------------------
# Tokenising and fingerprinting
# --------------------------------------------------------------------------

WORD_RE = re.compile(r"[A-Za-z0-9\u00C0-\u024F]+(?:['’][A-Za-z]+)?")


def tokenize(text):
    """Return list of (normalised_word, start_char, end_char)."""
    return [(m.group(0).lower().replace("’", "'"), m.start(), m.end()) for m in WORD_RE.finditer(text)]


def shingles(words, n):
    """Hash every n-word window -> list of (hash, start_index)."""
    out = []
    for i in range(len(words) - n + 1):
        h = hashlib.blake2b(" ".join(words[i : i + n]).encode(), digest_size=8).digest()
        out.append((h, i))
    return out


# --------------------------------------------------------------------------
# Exclusions (Turnitin-style filters)
# --------------------------------------------------------------------------

BIB_HEAD = re.compile(
    r"^\s*(references|bibliography|works cited|reference list|sources|literature cited)\s*:?\s*$",
    re.I | re.M,
)
QUOTE_RE = re.compile(r"[\"“”](.{3,800}?)[\"“”]", re.S)


def excluded_positions(text, tokens, exclude_quotes, exclude_bib):
    """Return set of token indices to ignore, plus reason labels for the report."""
    excl, reasons = set(), {}
    spans = []
    if exclude_quotes:
        spans += [(m.start(), m.end(), "quote") for m in QUOTE_RE.finditer(text)]
    if exclude_bib:
        heads = list(BIB_HEAD.finditer(text))
        if heads:
            spans.append((heads[-1].start(), len(text), "bibliography"))
    for s, e, label in spans:
        for i, (_, ts, te) in enumerate(tokens):
            if ts >= s and te <= e:
                excl.add(i)
                reasons[i] = label
    return excl, reasons


# --------------------------------------------------------------------------
# Sources: local folder, repository, web
# --------------------------------------------------------------------------


def load_folder(folder: Path, label_prefix=""):
    sources = []
    if not folder or not folder.exists():
        return sources
    for p in sorted(folder.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            try:
                t = extract_text(p)
            except Exception as ex:  # noqa: BLE001
                print(f"  ! could not read {p.name}: {ex}", file=sys.stderr)
                continue
            if t.strip():
                sources.append({"name": label_prefix + p.name, "kind": label_prefix.strip(": ") or "file",
                                "location": str(p), "text": t})
    return sources


def _http_get(url, headers=None, timeout=15):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 origincheck/1.0", **(headers or {})})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        ctype = r.headers.get("Content-Type", "")
        data = r.read(3_000_000)
    return ctype, data


def web_sources(text, max_queries, api_key, results_per_query=3):
    """Search distinctive sentences via the Brave Search API and fetch the hit pages."""
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if len(s.split()) >= 12]
    if not sentences:
        return []
    step = max(1, len(sentences) // max_queries)
    picked = sentences[::step][:max_queries]
    urls = []
    for i, s in enumerate(picked, 1):
        q = " ".join(s.split()[:24])
        print(f"  web query {i}/{len(picked)}", file=sys.stderr)
        try:
            _, data = _http_get(
                "https://api.search.brave.com/res/v1/web/search?count=%d&q=%s"
                % (results_per_query, urllib.parse.quote(f'"{q}"' if len(q) < 180 else q)),
                {"X-Subscription-Token": api_key, "Accept": "application/json"},
            )
            for r in json.loads(data).get("web", {}).get("results", []):
                if r.get("url") and r["url"] not in urls:
                    urls.append(r["url"])
        except Exception as ex:  # noqa: BLE001
            print(f"  ! search failed: {ex}", file=sys.stderr)
        time.sleep(1.1)  # free tier: 1 request/second
    out = []
    for u in urls:
        try:
            ctype, data = _http_get(u)
            if "pdf" in ctype or u.lower().endswith(".pdf"):
                import io
                from pypdf import PdfReader

                t = "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(data)).pages)
            else:
                t = html_to_text(data.decode("utf-8", errors="ignore"))
            if len(t.split()) > 50:
                out.append({"name": urllib.parse.urlparse(u).netloc + urllib.parse.urlparse(u).path[:60],
                            "kind": "internet", "location": u, "text": t})
        except Exception as ex:  # noqa: BLE001
            print(f"  ! fetch failed {u}: {ex}", file=sys.stderr)
    return out


# --------------------------------------------------------------------------
# Matching
# --------------------------------------------------------------------------


def find_matches(doc_tokens, sources, n, min_words, excluded):
    doc_words = [t[0] for t in doc_tokens]
    doc_sh = shingles(doc_words, n)

    # word index -> set(source ids) that cover it
    cover = defaultdict(set)
    for sid, src in enumerate(sources):
        stoks = tokenize(src["text"])
        src["_words"] = len(stoks)
        index = defaultdict(list)
        for h, i in shingles([t[0] for t in stoks], n):
            index[h].append(i)
        for h, i in doc_sh:
            if h in index:
                for k in range(i, i + n):
                    if k not in excluded:
                        cover[k].add(sid)

    # Build contiguous runs per source and drop runs shorter than min_words
    per_source = defaultdict(set)
    for sid in range(len(sources)):
        run = []
        for k in range(len(doc_words) + 1):
            if k < len(doc_words) and sid in cover.get(k, ()):
                run.append(k)
            else:
                if len(run) >= min_words:
                    per_source[sid].update(run)
                run = []

    # Turnitin-style attribution: each word credited to one source only,
    # preferring the source with the most matched words overall.
    ranking = sorted(per_source, key=lambda s: -len(per_source[s]))
    owner = {}
    for sid in ranking:
        for k in per_source[sid]:
            owner.setdefault(k, sid)
    return owner


# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------

PALETTE = ["#e5484d", "#3e63dd", "#30a46c", "#f76b15", "#8e4ec6", "#12a594", "#d6409f",
           "#ad7f58", "#0090ff", "#c2a300", "#e54666", "#46a758"]


def build_report(doc_name, text, tokens, owner, excluded, reasons, sources, settings):
    total = sum(1 for i in range(len(tokens)) if i not in excluded) or 1
    counts = defaultdict(int)
    for sid in owner.values():
        counts[sid] += 1
    order = sorted(counts, key=lambda s: -counts[s])
    num = {sid: i + 1 for i, sid in enumerate(order)}
    color = {sid: PALETTE[i % len(PALETTE)] for i, sid in enumerate(order)}
    index = round(100 * len(owner) / total)
    by_kind = defaultdict(int)
    for sid, c in counts.items():
        by_kind[sources[sid]["kind"]] += c

    # Highlighted body
    out, pos, i = [], 0, 0
    while i < len(tokens):
        sid = owner.get(i)
        if sid is None and i not in excluded:
            i += 1
            continue
        j = i
        while j + 1 < len(tokens) and (owner.get(j + 1) == sid if sid is not None else (j + 1) in excluded and owner.get(j + 1) is None):
            j += 1
        s, e = tokens[i][1], tokens[j][2]
        out.append(html.escape(text[pos:s]))
        frag = html.escape(text[s:e])
        if sid is not None:
            out.append(f'<mark style="--c:{color[sid]}" title="{html.escape(sources[sid]["name"])}">'
                       f'<sup>{num[sid]}</sup>{frag}</mark>')
        else:
            out.append(f'<span class="ex" title="Excluded: {reasons.get(i, "")}">{frag}</span>')
        pos, i = e, j + 1
    out.append(html.escape(text[pos:]))
    body = "".join(out).replace("\n\n", "</p><p>").replace("\n", "<br>")

    rows = []
    for sid in order:
        src = sources[sid]
        pct = 100 * counts[sid] / total
        loc = src["location"]
        link = f'<a href="{html.escape(loc)}">{html.escape(src["name"])}</a>' if loc.startswith("http") \
            else html.escape(src["name"])
        rows.append(f'<li><span class="n" style="background:{color[sid]}">{num[sid]}</span>'
                    f'<div><div class="sn">{link}</div><div class="sk">{html.escape(src["kind"])} · '
                    f'{counts[sid]} words</div></div><b>{pct:.0f}%</b></li>')
    kinds = " · ".join(f"{k.title()} {100 * v / total:.0f}%" for k, v in by_kind.items()) or "No matches"
    band = "low" if index < 15 else "mid" if index < 25 else "mid2" if index < 40 else "high"

    return f"""<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Similarity Report – {html.escape(doc_name)}</title>
<style>
:root{{--bg:#f6f7f9;--card:#fff;--ink:#1c2024;--mute:#60646c;--line:#e0e1e6}}
@media (prefers-color-scheme:dark){{:root{{--bg:#111113;--card:#18191b;--ink:#edeef0;--mute:#b0b4ba;--line:#2e3135}}}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--ink);font:15px/1.6 system-ui,Segoe UI,sans-serif}}
.wrap{{display:grid;grid-template-columns:minmax(0,1fr) 340px;gap:20px;max-width:1300px;margin:auto;padding:20px 16px}}
@media (max-width:900px){{.wrap{{grid-template-columns:1fr}}aside{{order:-1}}}}
.card{{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:20px 24px}}
h1{{font-size:18px;margin:0 0 4px}}.meta{{color:var(--mute);font-size:13px}}
.doc p{{margin:0 0 1em}}.doc{{font-family:Georgia,serif;font-size:16px}}
mark{{background:color-mix(in srgb,var(--c) 22%,transparent);color:inherit;border-bottom:2px solid var(--c);padding:0 1px;border-radius:2px}}
mark sup{{background:var(--c);color:#fff;font:600 10px system-ui;padding:0 4px;border-radius:3px;margin-right:2px}}
.ex{{color:var(--mute);text-decoration:line-through dotted;text-decoration-color:var(--line)}}
aside .card{{position:sticky;top:20px}}
.idx{{font-size:56px;font-weight:700;line-height:1}}.idx small{{font-size:22px}}
.low{{color:#30a46c}}.mid{{color:#c2a300}}.mid2{{color:#f76b15}}.high{{color:#e5484d}}
ol{{list-style:none;padding:0;margin:16px 0 0}}li{{display:flex;gap:10px;align-items:center;padding:8px 0;border-top:1px solid var(--line)}}
li div{{flex:1;min-width:0}}.sn{{overflow:hidden;text-overflow:ellipsis;white-space:nowrap;font-weight:500}}.sk{{color:var(--mute);font-size:12px}}
.n{{color:#fff;font-weight:700;font-size:12px;min-width:22px;text-align:center;border-radius:4px;padding:2px 4px}}
a{{color:inherit}}.set{{color:var(--mute);font-size:12px;margin-top:14px;border-top:1px solid var(--line);padding-top:10px}}
</style></head><body><div class="wrap">
<main class="card"><h1>{html.escape(doc_name)}</h1>
<div class="meta">{total} words checked · generated {datetime.now():%d %b %Y %H:%M}</div><hr style="border:0;border-top:1px solid var(--line);margin:14px 0">
<div class="doc"><p>{body}</p></div></main>
<aside><div class="card"><div class="meta">Similarity index</div>
<div class="idx {band}">{index}<small>%</small></div><div class="meta">{kinds}</div>
<ol>{''.join(rows) or '<li>No matching sources found.</li>'}</ol>
<div class="set">Settings: {html.escape(settings)}<br>Sources compared: {len(sources)}<br>
Strikethrough text was excluded from the check.</div></div></aside>
</div></body></html>"""


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser(description="Turnitin-style local similarity checker")
    ap.add_argument("document", help="submission to check (.docx .pdf .txt .md .html .rtf)")
    ap.add_argument("-s", "--sources", action="append", default=[], help="folder of source documents (repeatable)")
    ap.add_argument("-o", "--out", help="report path (default: <document>_similarity.html)")
    ap.add_argument("--ngram", type=int, default=5, help="words per fingerprint (default 5; lower = stricter)")
    ap.add_argument("--min-words", type=int, default=8, help="ignore matches shorter than this many words (default 8)")
    ap.add_argument("--exclude-quotes", action="store_true", help="ignore text inside quotation marks")
    ap.add_argument("--exclude-bib", action="store_true", help="ignore the References/Bibliography section")
    ap.add_argument("--exclude-source-pct", type=float, default=0,
                    help="drop sources contributing less than this %% (Turnitin 'exclude small sources')")
    ap.add_argument("--web", action="store_true", help="also search the internet (needs BRAVE_API_KEY env var)")
    ap.add_argument("--web-queries", type=int, default=15, help="max web searches (default 15)")
    ap.add_argument("--no-repo", action="store_true", help="don't compare against the saved-submissions repository")
    ap.add_argument("--save-to-repo", action="store_true", help="store this document in the repository after checking")
    ap.add_argument("--json", action="store_true", help="also write a machine-readable .json result")
    a = ap.parse_args()

    doc = Path(a.document)
    if not doc.exists() or doc.suffix.lower() not in SUPPORTED:
        sys.exit(f"Can't read {doc}. Supported: {', '.join(sorted(SUPPORTED))}")

    print(f"Reading {doc.name} ...", file=sys.stderr)
    text = extract_text(doc)
    tokens = tokenize(text)
    if len(tokens) < a.ngram:
        sys.exit("Document has too little text to check.")

    sources = []
    for f in a.sources:
        sources += load_folder(Path(f))
    if not a.no_repo:
        repo = [s for s in load_folder(REPO_DIR, "repository: ") if s["name"] != "repository: " + doc.name]
        sources += repo
    if a.web:
        key = os.environ.get("BRAVE_API_KEY")
        if not key:
            print("  ! --web needs a BRAVE_API_KEY environment variable; skipping web check.", file=sys.stderr)
        else:
            print("Searching the web ...", file=sys.stderr)
            sources += web_sources(text, a.web_queries, key)
    if not sources:
        sys.exit("No sources to compare against. Pass --sources <folder> and/or --web.")
    print(f"Comparing against {len(sources)} source(s) ...", file=sys.stderr)

    excluded, reasons = excluded_positions(text, tokens, a.exclude_quotes, a.exclude_bib)
    owner = find_matches(tokens, sources, a.ngram, a.min_words, excluded)

    total = sum(1 for i in range(len(tokens)) if i not in excluded) or 1
    if a.exclude_source_pct > 0:
        counts = defaultdict(int)
        for sid in owner.values():
            counts[sid] += 1
        small = {s for s, c in counts.items() if 100 * c / total < a.exclude_source_pct}
        owner = {k: v for k, v in owner.items() if v not in small}

    settings = (f"{a.ngram}-word fingerprints, min match {a.min_words} words"
                + (", quotes excluded" if a.exclude_quotes else "")
                + (", bibliography excluded" if a.exclude_bib else "")
                + (f", sources <{a.exclude_source_pct:g}% excluded" if a.exclude_source_pct else ""))
    report = build_report(doc.name, text, tokens, owner, excluded, reasons, sources, settings)
    out = Path(a.out) if a.out else doc.with_name(doc.stem + "_similarity.html")
    out.write_text(report, encoding="utf-8")

    index = round(100 * len(owner) / total)
    counts = defaultdict(int)
    for sid in owner.values():
        counts[sid] += 1
    print(f"\nSimilarity index: {index}%")
    for sid in sorted(counts, key=lambda s: -counts[s])[:10]:
        print(f"  {100 * counts[sid] / total:5.1f}%  {sources[sid]['name']}")
    print(f"\nReport: {out}")

    if a.json:
        jpath = out.with_suffix(".json")
        jpath.write_text(json.dumps({
            "document": doc.name, "similarity_index": index, "words_checked": total,
            "sources": [{"name": sources[s]["name"], "location": sources[s]["location"],
                         "percent": round(100 * counts[s] / total, 1)} for s in sorted(counts, key=lambda s: -counts[s])],
        }, indent=2), encoding="utf-8")
        print(f"JSON:   {jpath}")

    if a.save_to_repo:
        REPO_DIR.mkdir(exist_ok=True)
        (REPO_DIR / (doc.stem + ".txt")).write_text(text, encoding="utf-8")
        print(f"Saved to repository: {REPO_DIR / (doc.stem + '.txt')}")


if __name__ == "__main__":
    main()
